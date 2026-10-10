# -*- coding: utf-8 -*-
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("work/maritime-anomaly-agent/马六甲海峡AIS异常风险AHP评价标准.docx").resolve()
RI = {1: 0, 2: 0, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24, 7: 1.32, 8: 1.41, 9: 1.45}


def ahp(matrix):
    a = np.array(matrix, dtype=float)
    vals, vecs = np.linalg.eig(a)
    idx = np.argmax(vals.real)
    lam = float(vals[idx].real)
    weights = np.abs(vecs[:, idx].real)
    weights = weights / weights.sum()
    n = a.shape[0]
    ci = (lam - n) / (n - 1) if n > 1 else 0.0
    cr = ci / RI[n] if RI.get(n) else 0.0
    return weights, lam, ci, cr


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_text(cell, text, bold=False, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_text(t.rows[0].cells[i], h, True, font_size)
        shade(t.rows[0].cells[i], "E8EEF5")
        if widths:
            t.rows[0].cells[i].width = widths[i]
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_text(cells[i], val, False, font_size)
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return t


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.color.rgb = RGBColor(46, 116, 181) if level <= 2 else RGBColor(31, 77, 120)
    return p


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)


def fmt(x):
    return f"{x:.4f}"


def pct(x):
    return f"{x * 100:.2f} 分"


def matrix_rows(names, matrix):
    return [[names[i]] + [fmt(v) for v in row] for i, row in enumerate(matrix)]


def build():
    criteria_names = ["船舶行为异常", "碰撞与会遇风险", "水域通航环境", "船舶类型后果", "AIS数据可信度"]
    criteria_matrix = [
        [1, 1, 2, 3, 4],
        [1, 1, 2, 3, 4],
        [1 / 2, 1 / 2, 1, 2, 3],
        [1 / 3, 1 / 3, 1 / 2, 1, 2],
        [1 / 4, 1 / 4, 1 / 3, 1 / 2, 1],
    ]
    criteria_w, criteria_lam, criteria_ci, criteria_cr = ahp(criteria_matrix)

    secondary = {
        "船舶行为异常": (
            ["偏离常规航线", "异常停留", "航向突变", "航速异常"],
            [[1, 2, 3, 3], [1 / 2, 1, 2, 2], [1 / 3, 1 / 2, 1, 1], [1 / 3, 1 / 2, 1, 1]],
        ),
        "碰撞与会遇风险": (
            ["DCPA最近会遇距离", "TCPA最近会遇时间", "附近船舶密度"],
            [[1, 1, 2], [1, 1, 2], [1 / 2, 1 / 2, 1]],
        ),
        "水域通航环境": (
            ["主航道/港口附近", "狭窄或交汇水域", "敏感区域接近"],
            [[1, 1, 2], [1, 1, 2], [1 / 2, 1 / 2, 1]],
        ),
        "船舶类型后果": (
            ["船舶类别后果", "危险货物可能性"],
            [[1, 2], [1 / 2, 1]],
        ),
        "AIS数据可信度": (
            ["关键字段缺失", "数据更新间隔异常", "船型未知"],
            [[1, 1, 2], [1, 1, 2], [1 / 2, 1 / 2, 1]],
        ),
    }

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("马六甲海峡 AIS 异常风险 AHP 评价标准")
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(11, 37, 69)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("权重由 AHP 判断矩阵计算，并结合 AIS 异常检测、碰撞风险和马六甲海峡通航风险文献解释")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(99, 112, 131)

    heading(doc, "一、这套权重如何保证不是主观乱编", 1)
    para(doc, "本评价体系采用“文献确定指标、AHP 计算权重、一致性检验校验”的流程。也就是说，指标维度不是凭感觉列出来，权重也不是直接手填，而是通过两两比较矩阵计算得到。答辩时可以说明：判断矩阵中的重要性比较来自相关研究中反复强调的风险因素，并结合马六甲海峡船流密集、航道狭窄、油轮货船通行量大等场景特点。")
    bullet(doc, "AIS 异常检测研究常关注偏离常规航线、异常停留、航速异常、航向异常、AIS 数据异常等行为。")
    bullet(doc, "船舶碰撞风险研究常使用 DCPA、TCPA、会遇态势、附近船舶密度等指标。")
    bullet(doc, "马六甲海峡相关研究强调该水域船流密集、航道复杂、碰撞和拥堵风险突出，适合用 AIS 和 AHP 做航行安全评价。")
    bullet(doc, "AHP 权重通过判断矩阵计算，并使用 CR 一致性比例检验；通常 CR < 0.10 认为判断矩阵可以接受。")

    heading(doc, "二、文献依据与指标对应关系", 1)
    table(
        doc,
        ["文献/资料方向", "对本评价体系的支持", "对应维度"],
        [
            ["IMO AIS 资料", "AIS 可提供船舶身份、类型、位置、航向、速度、航行状态等信息，是监测和评价的基础。", "AIS 数据可信度、船舶行为异常"],
            ["马六甲海峡 AIS + AHP 航行安全研究", "说明马六甲海峡适合用 AIS 数据和 AHP 方法开展航行安全评价。", "水域通航环境、碰撞与会遇风险"],
            ["AIS 异常检测综述", "总结了偏离航线、异常航速/航向、异常停留、AIS 异常活动等常见异常。", "船舶行为异常、AIS 数据可信度"],
            ["船舶碰撞风险研究", "DCPA、TCPA、交通密度是衡量船舶会遇和碰撞风险的常用指标。", "碰撞与会遇风险"],
            ["马六甲高密度水域风险研究", "高密度船流、港口附近、航道交汇和狭窄水域会放大异常后果。", "水域通航环境、船舶类型后果"],
        ],
        widths=[Cm(4), Cm(8), Cm(4)],
        font_size=8.5,
    )

    heading(doc, "三、一级维度 AHP 判断矩阵与权重", 1)
    para(doc, "一级维度保留 5 项风险因素，去掉上一版中的“智能体解释与处置价值”。原因是解释能力属于系统功能表现，不是海上风险本身；把它放入风险总分，答辩时容易被质疑。")
    table(
        doc,
        ["维度"] + criteria_names,
        matrix_rows(criteria_names, criteria_matrix),
        widths=[Cm(3)] + [Cm(2.4)] * 5,
        font_size=7.5,
    )
    table(
        doc,
        ["一级维度", "AHP 权重", "折算分值", "权重解释"],
        [
            ["船舶行为异常", fmt(criteria_w[0]), pct(criteria_w[0]), "AIS 异常检测首先判断船舶自身轨迹、速度、航向等行为是否异常。"],
            ["碰撞与会遇风险", fmt(criteria_w[1]), pct(criteria_w[1]), "马六甲海峡船舶密集，DCPA/TCPA 等会遇指标直接关系航行安全。"],
            ["水域通航环境", fmt(criteria_w[2]), pct(criteria_w[2]), "主航道、港口、狭窄水域会放大同一异常行为的风险后果。"],
            ["船舶类型后果", fmt(criteria_w[3]), pct(criteria_w[3]), "油轮、货船等不同船型发生异常后的后果严重程度不同。"],
            ["AIS数据可信度", fmt(criteria_w[4]), pct(criteria_w[4]), "数据缺失、更新异常、船型未知会影响判断可靠性，但通常不是风险本体。"],
        ],
        widths=[Cm(3.2), Cm(2), Cm(2.2), Cm(8.2)],
        font_size=8.5,
    )
    para(doc, f"一级判断矩阵计算结果：λmax={criteria_lam:.4f}，CI={criteria_ci:.4f}，CR={criteria_cr:.4f}。CR < 0.10，因此一致性检验通过。")

    heading(doc, "四、一级维度两两比较理由", 1)
    table(
        doc,
        ["比较关系", "判断值", "依据说明"],
        [
            ["船舶行为异常 与 碰撞与会遇风险", "1", "二者同等重要：前者识别异常来源，后者判断是否会直接造成航行危险。"],
            ["船舶行为异常 / 水域通航环境", "2", "异常监测首先看船舶自身行为，水域环境主要起风险放大作用。"],
            ["船舶行为异常 / 船舶类型后果", "3", "船型会影响后果，但没有异常行为时通常不会单独触发预警。"],
            ["船舶行为异常 / AIS数据可信度", "4", "数据可信度影响判断质量，但行为异常是风险识别的主体。"],
            ["碰撞与会遇风险 / 水域通航环境", "2", "马六甲场景下，短时间会遇风险比一般环境条件更直接。"],
            ["碰撞与会遇风险 / 船舶类型后果", "3", "DCPA/TCPA 等指标直接对应碰撞风险，船型更多影响后果严重性。"],
            ["水域通航环境 / 船舶类型后果", "2", "同一船型在主航道、港口或狭窄水域内风险会明显升高。"],
            ["船舶类型后果 / AIS数据可信度", "2", "油轮/货船等后果差异对预警优先级较明显，数据可信度主要用于复核。"],
        ],
        widths=[Cm(4.4), Cm(1.6), Cm(9.6)],
        font_size=8.5,
    )

    heading(doc, "五、二级指标 AHP 权重与最终 100 分折算", 1)
    final_rows = []
    cr_rows = []
    for i, cname in enumerate(criteria_names):
        names, mat = secondary[cname]
        w, lam, ci, cr = ahp(mat)
        cr_rows.append([cname, f"{lam:.4f}", f"{ci:.4f}", f"{cr:.4f}", "通过" if cr < 0.10 else "需调整"])
        for j, sname in enumerate(names):
            final = float(criteria_w[i] * w[j])
            final_rows.append([cname, fmt(criteria_w[i]), sname, fmt(w[j]), fmt(final), pct(final)])
    table(
        doc,
        ["一级维度", "一级权重", "二级指标", "二级局部权重", "最终权重", "100分折算"],
        final_rows,
        widths=[Cm(2.7), Cm(1.6), Cm(3.2), Cm(1.8), Cm(1.7), Cm(1.9)],
        font_size=7.8,
    )
    table(
        doc,
        ["二级指标所属维度", "λmax", "CI", "CR", "一致性结果"],
        cr_rows,
        widths=[Cm(4), Cm(2.2), Cm(2.2), Cm(2.2), Cm(3)],
        font_size=8.5,
    )
    para(doc, "所有二级指标最终权重相加为 1，对应总分 100 分。实际运行时，每个二级指标先按规则得到 0-100 分，再乘以最终权重，最后求和。")

    heading(doc, "六、二级指标怎么打 0-100 分", 1)
    table(
        doc,
        ["二级指标", "低分情况", "中等情况", "高分情况", "建议计算方法"],
        [
            ["偏离常规航线", "位于常规航道内", "轻微偏离航道", "明显偏离航道或进入不常见区域", "计算点到主航道中心线/历史轨迹带的距离。"],
            ["异常停留", "正常航行", "低速或静止 10-30 分钟", "低速或静止超过 30 分钟且处于敏感水域", "先设 SOG < 1 节为低速阈值，按持续时间评分。"],
            ["航向突变", "航向平稳", "短时间变化 30-60 度", "短时间变化超过 60 度", "计算相邻 AIS 点 COG/heading 差值。"],
            ["航速异常", "速度平稳", "速度突然变化 5-10 节", "速度突变超过 10 节或长期异常低速", "计算相邻点 SOG 差值和低速持续时间。"],
            ["DCPA最近会遇距离", "最近距离充足", "最近距离偏小", "最近距离很小，存在碰撞风险", "预测两船最近距离，距离越小分越高。"],
            ["TCPA最近会遇时间", "短期内不会会遇", "10-30 分钟内会遇", "10 分钟内会遇", "预测到达最近点时间，时间越短分越高。"],
            ["附近船舶密度", "周围船少", "周围船舶较多", "周围船舶密集", "统计 1-3 海里范围内船舶数量。"],
            ["主航道/港口附近", "远离繁忙区域", "靠近航道边缘", "位于主航道、港口、锚地附近", "根据与重点区域距离评分。"],
            ["狭窄或交汇水域", "开阔水域", "一般受限水域", "狭窄、转弯或航道交汇水域", "预先划定区域风险等级。"],
            ["敏感区域接近", "未接近", "接近边缘", "进入敏感区域", "按与禁航区、锚地、港区等距离评分。"],
            ["船舶类别后果", "低后果船型", "货船等中等后果船型", "油轮/危险品相关船型", "结合 AIS vessel_type 和分类结果赋分。"],
            ["危险货物可能性", "无明显危险货物", "疑似相关", "高度可能运输油品/危险品", "结合船型、船名、目的地、人工补全表判断。"],
            ["关键字段缺失", "字段完整", "部分字段缺失", "经纬度/速度/航向等关键字段缺失", "统计 lat/lon/sog/cog/heading/type 缺失比例。"],
            ["数据更新间隔异常", "更新连续", "更新间隔偏长", "长时间无更新或突然消失", "计算同一 MMSI 相邻时间间隔。"],
            ["船型未知", "船型明确", "船型不确定", "无法自动分类且人工表无结果", "未知船提高复核风险。"],
        ],
        widths=[Cm(3), Cm(2.6), Cm(3.1), Cm(3.5), Cm(4.1)],
        font_size=7.7,
    )

    heading(doc, "七、最终总分计算方法", 1)
    para(doc, "总分 = Σ（二级指标得分 × 二级指标最终权重）。每个二级指标得分取 0-100，所有最终权重相加为 1，因此总分也是 0-100 分。")
    para(doc, "也可以分两步算：第一步，一级维度分 = Σ（二级指标得分 × 二级局部权重）；第二步，最终总分 = Σ（一级维度分 × 一级权重）。两种算法结果一致。")
    para(doc, "举例：如果某异常事件“偏离常规航线”为 80 分、“异常停留”为 60 分、“航向突变”为 20 分、“航速异常”为 40 分，则船舶行为异常维度得分 = 80×0.4554 + 60×0.2628 + 20×0.1409 + 40×0.1409 = 60.65 分。再乘以一级权重 0.3192，该维度对总分贡献约 19.36 分。其他维度同理计算，最后相加得到总风险分。")
    table(
        doc,
        ["总分区间", "风险等级", "Demo 展示建议"],
        [["0-25", "低风险", "正常观察"], ["26-50", "中风险", "需要关注"], ["51-75", "高风险", "建议预警"], ["76-100", "严重风险", "重点人工复核"]],
        widths=[Cm(3), Cm(3), Cm(8)],
        font_size=8.5,
    )

    heading(doc, "八、参考资料", 1)
    refs = [
        "IMO AIS: https://www.imo.org/en/ourwork/safety/pages/ais.aspx",
        "Implementation of AIS for Evaluation of Marine Traffic Safety in Strait of Malacca using AHP: https://www.jstage.jst.go.jp/article/jjasnaoe/16/0/16_141/_article",
        "Risk of Navigation for Marine Traffic in the Malacca Strait Using AIS: https://www.sciencedirect.com/science/article/pii/S1878522015002337",
        "Anomaly Detection in Maritime AIS Tracks: A Review of Recent Approaches: https://www.mdpi.com/2077-1312/10/1/112",
        "New frontiers in the risk assessment of ship collision: https://doi.org/10.1016/J.OCEANENG.2023.113999",
        "AHPy GitHub 项目，可参考其 AHP 权重和一致性检验实现: https://github.com/PhilipGriffith/AHPy",
    ]
    for ref in refs:
        bullet(doc, ref)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print("criteria_cr", criteria_cr)
    print("final_weight_sum", sum(float(r[4]) for r in final_rows))


if __name__ == "__main__":
    build()
